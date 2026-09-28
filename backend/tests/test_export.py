"""
Comprehensive Unit and Integration Tests for Exporter Engine and Transaction Workspace.
Tests:
- RFC 4180 CSV export with UTF-8 BOM, comma/quote escaping, valid header structure.
- Formatted Excel (.xlsx) export with openpyxl, dark navy styling, numeric float values, formulas, sheets.
- Versioned structured JSON export.
- Programmatic round-trip lossless verification Ingest(Export(S)) == S.
- Export REST API endpoints (/api/export/csv, /xlsx, /json, /verify).
- Transaction update and reconciliation endpoints (/api/statements/{id}/transactions, /api/statements/reconcile).
"""
import io
import json
import csv
from decimal import Decimal
import pytest
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.schemas.statement import StatementMetadata
from app.schemas.transaction import TransactionRecord
from app.schemas.reconciliation import ReconciliationSummary
from app.services.reconciliation import to_decimal, reconcile_statement
from app.services.exporter import (
    ExportPayload,
    export_to_csv,
    export_to_xlsx,
    export_to_json,
    verify_round_trip,
)


@pytest.fixture
def sample_payload():
    meta = StatementMetadata(
        bank_name="First National Commercial Bank",
        account_number="**** **** **** 5678",
        statement_period_start="2026-08-01",
        statement_period_end="2026-08-31",
        starting_balance="2500.00",
        ending_balance="4250.75",
        currency="USD",
        page_count=2,
    )
    txs = [
        TransactionRecord(
            id="tx_01",
            date="2026-08-02",
            payee="Stripe Payout, Inc.",
            type="credit",
            amount="2000.00",
            category="Revenue",
            running_balance="4500.00",
        ),
        TransactionRecord(
            id="tx_02",
            date="2026-08-05",
            payee='Office Supplies "Special" Order',
            type="debit",
            amount="149.25",
            category="Office",
            running_balance="4350.75",
        ),
        TransactionRecord(
            id="tx_03",
            date="2026-08-10",
            payee="Cloud Hosting, Server A & B",
            type="debit",
            amount="100.00",
            category="Software",
            running_balance="4250.75",
        ),
    ]
    rec_summary, _ = reconcile_statement("2500.00", "4250.75", txs)
    return ExportPayload(metadata=meta, reconciliation=rec_summary, transactions=txs)


# ===========================================================================
# 1. CSV Exporter Tests
# ===========================================================================

def test_csv_exporter_utf8_bom(sample_payload):
    csv_bytes = export_to_csv(sample_payload)
    assert csv_bytes.startswith(b"\xef\xbb\xbf"), "Missing mandatory UTF-8 BOM"


def test_csv_exporter_headers_and_metadata(sample_payload):
    csv_bytes = export_to_csv(sample_payload)
    text = csv_bytes.decode("utf-8-sig")
    lines = text.split("\r\n")

    assert "# Bank Statement Export,First National Commercial Bank" in lines[0]
    assert "# Account,**** **** **** 5678" in lines[1]
    assert "# Period,2026-08-01 to 2026-08-31" in lines[2]
    assert "# Starting Balance,2500.00" in lines[3]
    assert "# Ending Balance,4250.75" in lines[4]
    assert "# Reconciled,True" in lines[5]

    # Find table header
    table_header_idx = -1
    for i, line in enumerate(lines):
        if "Date,Description / Payee,Type,Amount,Category,Running Balance" in line:
            table_header_idx = i
            break
    assert table_header_idx != -1, "Transaction table header not found"


def test_csv_exporter_escapes_commas_and_quotes(sample_payload):
    csv_bytes = export_to_csv(sample_payload)
    text = csv_bytes.decode("utf-8-sig")

    # Double quotes escaping: "Special" -> ""Special""
    assert '""Special""' in text
    # Comma escaping: "Stripe Payout, Inc."
    assert '"Stripe Payout, Inc."' in text


def test_csv_exporter_row_counts_and_amounts(sample_payload):
    csv_bytes = export_to_csv(sample_payload)
    text = csv_bytes.decode("utf-8-sig")
    reader = list(csv.reader(io.StringIO(text)))

    tx_start = -1
    for idx, row in enumerate(reader):
        if row and "Date" in row[0]:
            tx_start = idx + 1
            break

    data_rows = reader[tx_start:]
    assert len(data_rows) == len(sample_payload.transactions)
    assert data_rows[0][3] == "2000.00"
    assert data_rows[1][3] == "149.25"
    assert data_rows[2][3] == "100.00"


# ===========================================================================
# 2. Excel (.xlsx) Exporter Tests
# ===========================================================================

def test_xlsx_exporter_loads_workbook(sample_payload):
    xlsx_bytes = export_to_xlsx(sample_payload)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    assert "Transactions" in wb.sheetnames
    assert "Reconciliation Summary" in wb.sheetnames


def test_xlsx_exporter_numeric_types_and_formatting(sample_payload):
    xlsx_bytes = export_to_xlsx(sample_payload)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    ws = wb["Transactions"]

    # Header check
    headers = [cell.value for cell in ws[1]]
    assert "Date" in headers
    assert "Amount" in headers
    assert "Running Balance" in headers

    # Amount in row 2 (Col 4) must be numeric
    amt_cell = ws.cell(row=2, column=4)
    assert isinstance(amt_cell.value, (int, float))
    assert amt_cell.value == 2000.0
    assert amt_cell.number_format == "$#,##0.00"

    # Running balance in row 2 (Col 6)
    bal_cell = ws.cell(row=2, column=6)
    assert isinstance(bal_cell.value, (int, float))
    assert bal_cell.value == 4500.0


def test_xlsx_exporter_header_styling(sample_payload):
    xlsx_bytes = export_to_xlsx(sample_payload)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    ws = wb["Transactions"]

    # Header styling
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.bold is True
    # Navy fill color
    assert header_cell.fill.start_color.rgb == "001E293B" or "1E293B" in str(header_cell.fill.start_color.rgb)


def test_xlsx_exporter_formulas_present(sample_payload):
    xlsx_bytes = export_to_xlsx(sample_payload)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    ws_tx = wb["Transactions"]

    # Find formula cells in column 4 (Amount)
    formulas = [cell.value for cell in ws_tx["D"] if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert any("SUMIF" in f for f in formulas), f"Expected SUMIF formulas, got {formulas}"

    # Check Summary sheet formulas
    ws_rec = wb["Reconciliation Summary"]
    calc_end_formula = ws_rec["B12"].value
    assert isinstance(calc_end_formula, str) and "=B9+B10-B11" in calc_end_formula
    disc_formula = ws_rec["B14"].value
    assert isinstance(disc_formula, str) and "=B12-B13" in disc_formula


# ===========================================================================
# 3. JSON Exporter Tests
# ===========================================================================

def test_json_exporter_structure_and_schema(sample_payload):
    json_bytes = export_to_json(sample_payload)
    data = json.loads(json_bytes.decode("utf-8"))

    assert "$schema" in data
    assert data["version"] == "1.0.0"
    assert "metadata" in data
    assert "reconciliation" in data
    assert "transactions" in data
    assert "category_summary" in data

    assert len(data["transactions"]) == len(sample_payload.transactions)
    assert data["reconciliation"]["is_reconciled"] is True
    assert data["metadata"]["starting_balance"] == "2500.00"
    assert data["metadata"]["ending_balance"] == "4250.75"


def test_json_exporter_category_summary(sample_payload):
    json_bytes = export_to_json(sample_payload)
    data = json.loads(json_bytes.decode("utf-8"))
    cats = {item["category"]: item for item in data["category_summary"]}

    assert "Revenue" in cats
    assert cats["Revenue"]["credits"] == "2000.00"
    assert "Office" in cats
    assert cats["Office"]["debits"] == "149.25"


# ===========================================================================
# 4. Programmatic Round-Trip Ingestion Tests
# ===========================================================================

def test_verify_round_trip_json(sample_payload):
    json_bytes = export_to_json(sample_payload)
    valid, msg = verify_round_trip(json_bytes, "json", sample_payload)
    assert valid is True
    assert "passed" in msg


def test_verify_round_trip_csv(sample_payload):
    csv_bytes = export_to_csv(sample_payload)
    valid, msg = verify_round_trip(csv_bytes, "csv", sample_payload)
    assert valid is True
    assert "passed" in msg


def test_verify_round_trip_xlsx(sample_payload):
    xlsx_bytes = export_to_xlsx(sample_payload)
    valid, msg = verify_round_trip(xlsx_bytes, "xlsx", sample_payload)
    assert valid is True
    assert "passed" in msg


def test_verify_round_trip_detects_tampered_json(sample_payload):
    json_bytes = export_to_json(sample_payload)
    data = json.loads(json_bytes.decode("utf-8"))
    # Corrupt a transaction amount
    data["transactions"][0]["amount"] = "99999.00"
    tampered_bytes = json.dumps(data).encode("utf-8")

    valid, msg = verify_round_trip(tampered_bytes, "json", sample_payload)
    assert valid is False
    assert "mismatch" in msg


# ===========================================================================
# 5. Export REST API Endpoint Tests
# ===========================================================================

def test_api_export_direct_csv(sample_payload):
    client = TestClient(app)
    response = client.post("/api/export/csv", json=sample_payload.model_dump())
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.content.startswith(b"\xef\xbb\xbf")


def test_api_export_direct_xlsx(sample_payload):
    client = TestClient(app)
    response = client.post("/api/export/xlsx", json=sample_payload.model_dump())
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Transactions" in wb.sheetnames


def test_api_export_direct_json(sample_payload):
    client = TestClient(app)
    response = client.post("/api/export/json", json=sample_payload.model_dump())
    assert response.status_code == 200
    data = response.json()
    assert data["version"] == "1.0.0"
    assert len(data["transactions"]) == len(sample_payload.transactions)


def test_api_export_verify_endpoint(sample_payload):
    client = TestClient(app)
    res = client.post("/api/export/verify/csv", json=sample_payload.model_dump())
    assert res.status_code == 200
    assert res.json()["valid"] is True


# ===========================================================================
# 6. Transaction Workspace API Tests
# ===========================================================================

def test_api_direct_reconcile():
    client = TestClient(app)
    txs = [
        {"id": "1", "date": "2026-08-01", "payee": "Deposit", "type": "credit", "amount": "1000.00", "running_balance": "2000.00", "category": "Income"},
        {"id": "2", "date": "2026-08-02", "payee": "Utility", "type": "debit", "amount": "250.00", "running_balance": "1750.00", "category": "Utilities"},
    ]
    payload = {
        "starting_balance": "1000.00",
        "reported_ending_balance": "1750.00",
        "transactions": txs,
    }
    res = client.post("/api/statements/reconcile", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["net_cashflow"] == "750.00"
    assert data["calculated_ending_balance"] == "1750.00"
    assert data["discrepancy"] == "0.00"
    assert data["is_reconciled"] is True


def test_api_statement_transaction_workflow():
    from app.core.database import SessionLocal
    from app.models.statement import StatementRecord
    from app.models.tenant import Tenant
    from app.models.user import User
    import uuid

    db = SessionLocal()
    try:
        # Create test tenant and statement
        tenant_id = f"ten_test_{uuid.uuid4().hex[:6]}"
        user_id = f"usr_test_{uuid.uuid4().hex[:6]}"
        stmt_id = f"stmt_test_{uuid.uuid4().hex[:6]}"

        tenant = Tenant(id=tenant_id, name="Test Tenant")
        user = User(id=user_id, tenant_id=tenant_id, email=f"{user_id}@test.com", hashed_password="dummy")
        stmt = StatementRecord(
            id=stmt_id,
            tenant_id=tenant_id,
            user_id=user_id,
            filename="sample_test.pdf",
            file_path="nonexistent.pdf",
            status="completed",
            page_count=1,
            starting_balance="500.00",
            ending_balance="700.00",
        )
        db.add(tenant)
        db.add(user)
        db.add(stmt)
        db.commit()

        client = TestClient(app)

        # 1. Update transactions via PUT
        new_txs = [
            {"id": "tx_a", "date": "2026-08-01", "payee": "Client Fee", "type": "credit", "amount": "300.00", "running_balance": "800.00", "category": "Revenue"},
            {"id": "tx_b", "date": "2026-08-02", "payee": "SaaS Subscription", "type": "debit", "amount": "100.00", "running_balance": "700.00", "category": "Software"},
        ]
        put_res = client.put(f"/api/statements/{stmt_id}/transactions", json={"transactions": new_txs})
        assert put_res.status_code == 200
        put_data = put_res.json()
        assert put_data["success"] is True
        assert put_data["reconciliation"]["is_reconciled"] is True
        assert put_data["reconciliation"]["net_cashflow"] == "200.00"

        # 2. Retrieve transactions via GET
        get_res = client.get(f"/api/statements/{stmt_id}/transactions")
        assert get_res.status_code == 200
        get_data = get_res.json()
        assert len(get_data["transactions"]) == 2
        assert get_data["reconciliation"]["is_reconciled"] is True

        # 3. Explicit re-reconciliation endpoint
        rec_res = client.post(f"/api/statements/{stmt_id}/reconcile")
        assert rec_res.status_code == 200
        assert rec_res.json()["is_reconciled"] is True

        # 4. Export by statement ID
        exp_res = client.get(f"/api/export/{stmt_id}/csv")
        assert exp_res.status_code == 200
        assert exp_res.content.startswith(b"\xef\xbb\xbf")
        assert b"Client Fee" in exp_res.content

        exp_xlsx = client.get(f"/api/export/{stmt_id}/xlsx")
        assert exp_xlsx.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(exp_xlsx.content))
        assert "Transactions" in wb.sheetnames

        exp_json = client.get(f"/api/export/{stmt_id}/json")
        assert exp_json.status_code == 200
        assert exp_json.json()["metadata"]["starting_balance"] == "500.00"

    finally:
        db.close()


def test_frontend_static_serving():
    client = TestClient(app)
    # Test index.html serving at root
    res = client.get("/")
    assert res.status_code == 200
    assert "html" in res.headers["content-type"]
    assert "StatementParser" in res.text or "<div id=\"root\">" in res.text

